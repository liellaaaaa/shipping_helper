# -*- coding: utf-8 -*-
"""
运输鉴定报告字段提取器
支持PDF和DOCX格式的运输鉴定报告
"""

import os
import re
from typing import Dict, Optional, List
import fitz  # PyMuPDF


class ReportExtractor:
    """运输鉴定报告字段提取器"""

    def __init__(self):
        self.fields = {}

    def extract_from_file(self, filepath: str) -> Dict[str, str]:
        """
        从文件提取字段

        Args:
            filepath: 文件路径（PDF或DOCX）

        Returns:
            提取的字段字典
        """
        ext = os.path.splitext(filepath)[1].lower()

        if ext == '.pdf':
            return self.extract_from_pdf(filepath)
        elif ext in ['.docx', '.doc']:
            return self.extract_from_docx(filepath)
        else:
            raise ValueError(f"不支持的文件格式: {ext}")

    def extract_from_pdf(self, filepath: str) -> Dict[str, str]:
        """从PDF提取字段"""
        self.fields = {}

        try:
            doc = fitz.open(filepath)
            full_text = ""

            for page in doc:
                full_text += page.get_text() + "\n"

            doc.close()

            self._extract_fields_from_text(full_text)
            return self.fields

        except Exception as e:
            raise RuntimeError(f"PDF提取失败: {str(e)}")

    def extract_from_docx(self, filepath: str) -> Dict[str, str]:
        """从DOCX提取字段"""
        from docx import Document

        self.fields = {}

        try:
            doc = Document(filepath)
            full_text = ""

            # 提取段落文本
            for para in doc.paragraphs:
                full_text += para.text + "\n"

            # 提取表格文本
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        full_text += cell.text + " "
                full_text += "\n"

            self._extract_fields_from_text(full_text)
            return self.fields

        except Exception as e:
            raise RuntimeError(f"DOCX提取失败: {str(e)}")

    def _extract_fields_from_text(self, text: str):
        """从文本中提取字段"""
        # 规范化文本：处理换行和多余空格
        lines = text.split('\n')
        normalized_lines = []
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            # 跳过纯空白行，但保留内容行
            if line:
                normalized_lines.append(line)
            i += 1

        # 方法：逐行分析，识别字段名和值
        n = len(normalized_lines)
        i = 0
        while i < n:
            line = normalized_lines[i]

            # 样品名称（中文）
            if '样品名称' in line and 'Name' not in line:
                # 向下查找实际内容（跳过标签行和英文名行）
                for j in range(i + 1, min(i + 4, n)):
                    content = normalized_lines[j]
                    # 跳过Name of Goods行和空行
                    if 'Name of Goods' in content or not content:
                        continue
                    # 如果包含中文，认为是实际的中文名称
                    if re.search(r'[一-鿿]', content):
                        self.fields['样品名称中文'] = content.strip()
                        break

            # 样品名称（英文）
            if 'Name of Goods' in line or (i > 0 and 'Name' in normalized_lines[i-1] and 'Name of Goods' in line):
                for j in range(i + 1, min(i + 4, n)):
                    content = normalized_lines[j]
                    if content and not re.search(r'[一-鿿]', content):
                        # 英文字符行
                        if re.search(r'[a-zA-Z]', content):
                            self.fields['样品名称英文'] = content.strip()
                            break

            # 报告编号 - 格式：No.XXXXXXXX
            match = re.search(r'No\.?\s*([A-Z0-9]{10,})', line)
            if match:
                self.fields['报告编号'] = match.group(1).strip()

            # 样品编号 - 格式：WXST2513478-036（在同一行或下一行）
            if 'Sample No' in line or '样品编号' in line:
                # 如果当前行就是编号值（没有其他标签）
                match = re.search(r'([A-Z0-9]{8,}-?\d*)', line)
                if match:
                    self.fields['样品编号'] = match.group(1).strip()
                else:
                    # 向下查找内容
                    for j in range(i + 1, min(i + 3, n)):
                        content = normalized_lines[j].strip()
                        if content:
                            match = re.search(r'([A-Z0-9]{8,}-?\d*)', content)
                            if match:
                                self.fields['样品编号'] = match.group(1).strip()
                                break

            # 样品描述（外观性状）- 可能跨行
            if '样品描述' in line:
                # 向下查找多行描述（跳过Appearance行）
                desc_lines = []
                skip_next = False  # 跳过紧跟的Appearance行
                for j in range(i + 1, min(i + 6, n)):
                    content = normalized_lines[j].strip()
                    if 'Appearance' in content:
                        skip_next = True
                        continue
                    if skip_next:
                        skip_next = False
                        continue
                    # 遇到鉴定结论就停止
                    if '鉴定结论' in content or 'Conclusion' in content:
                        break
                    if content:
                        desc_lines.append(content)
                    if len(desc_lines) >= 3:  # 可能有多行
                        break

                if desc_lines:
                    # 分割中英文（包含中文的为中文描述）
                    chinese = ''
                    english = ''
                    for line_content in desc_lines:
                        if re.search(r'[一-鿿]', line_content) and '鉴定结论' not in line_content:
                            chinese += line_content + ' '
                        elif re.search(r'[a-zA-Z]', line_content):
                            english += line_content + ' '

                    if chinese:
                        self.fields['外观性状中文'] = chinese.strip()
                    if english:
                        self.fields['外观性状英文'] = english.strip()

            # 委托单位
            if '委托单位' in line and 'Applicant' not in line:
                for j in range(i + 1, min(i + 4, n)):
                    content = normalized_lines[j].strip()
                    # 跳过纯翻译行（只有英文）
                    if content and not re.match(r'^[A-Za-z\s\.,]+$', content):
                        self.fields['委托单位'] = content
                        break

            # 生产单位
            if '生产单位' in line and 'Manufacturer' not in line:
                for j in range(i + 1, min(i + 4, n)):
                    content = normalized_lines[j].strip()
                    # 跳过纯翻译行（只有英文）
                    if content and not re.match(r'^[A-Za-z\s\.,]+$', content):
                        self.fields['生产单位'] = content
                        break

            # 鉴定结论
            if '鉴定结论' in line and 'Conclusion' in line:
                for j in range(i + 1, min(i + 4, n)):
                    content = normalized_lines[j].strip()
                    if content:
                        self.fields['鉴定结论'] = content
                        break

            i += 1

        # 提取运输类型（通过鉴定依据判断）
        text_lower = text.lower()
        if 'imo imdg' in text_lower or 'imdg code' in text_lower:
            self.fields['运输类型'] = '海运'
        elif 'iata dgr' in text_lower or '空运' in text:
            self.fields['运输类型'] = '空运'
        elif 'jt/t' in text_lower or '道路运输' in text:
            self.fields['运输类型'] = '陆运'

    def get_field(self, name: str) -> Optional[str]:
        """获取指定字段值"""
        return self.fields.get(name)

    def get_all_fields(self) -> Dict[str, str]:
        """获取所有提取的字段"""
        return self.fields.copy()


def extract_fields_from_reports_directory(directory: str) -> List[Dict[str, str]]:
    """
    从目录批量提取运输鉴定报告字段

    Args:
        directory: 目录路径

    Returns:
        每个文件的提取结果列表
    """
    results = []

    for filename in os.listdir(directory):
        filepath = os.path.join(directory, filename)
        if os.path.isfile(filepath):
            ext = os.path.splitext(filename)[1].lower()
            if ext in ['.pdf', '.docx', '.doc']:
                try:
                    extractor = ReportExtractor()
                    fields = extractor.extract_from_file(filepath)
                    fields['_filename'] = filename
                    fields['_filepath'] = filepath
                    results.append(fields)
                except Exception as e:
                    print(f"提取失败 {filename}: {e}")

    return results