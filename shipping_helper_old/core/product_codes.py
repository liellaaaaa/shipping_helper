# -*- coding: utf-8 -*-
"""
商品编码表加载器
从Excel文件加载产品编码、报关成分等信息
"""

import os
import openpyxl
from typing import Dict, List, Optional


class ProductCodesLoader:
    """商品编码表加载器"""

    # 列索引（从0开始）
    COL_INDEX = 0
    COL_PRODUCT_CODE = 1  # 产品内编
    COL_PI_NAME = 2  # PI名建议
    COL_PRODUCT_LINE = 3  # 产品线
    COL_APPEARANCE = 4  # 产品外观
    COL_OLD_CODE = 5  # 旧商品编码
    COL_NEW_CODE = 6  # 新商品编码
    COL_DECLARE_NAME = 7  # 报关名称
    COL_TRANSPORT_NAME = 8  # 运输鉴定名称
    COL_USAGE = 9  # 用途
    COL_COMPONENTS = 10  # 成分
    COL_TAX_REFUND = 11  # 退税税率
    COL_REGULATION = 12  # 监管条件
    COL_INSPECTION = 13  # 检验检疫
    # 后面的列暂时忽略

    def __init__(self):
        self.filepath = None
        self.workbook = None
        self.worksheet = None
        self.products = []  # 产品列表
        self.products_dict = {}  # 产品内编 -> 产品信息

    def load(self, filepath: str) -> bool:
        """
        加载商品编码Excel文件

        Args:
            filepath: Excel文件路径

        Returns:
            是否加载成功
        """
        if not os.path.exists(filepath):
            return False

        try:
            self.filepath = filepath
            self.workbook = openpyxl.load_workbook(filepath, data_only=True)
            self.worksheet = self.workbook.active

            # 读取所有产品数据
            self._load_products()
            return True

        except Exception as e:
            print(f"加载商品编码表失败: {e}")
            return False

    def _load_products(self):
        """从Excel加载产品数据"""
        self.products = []
        self.products_dict = {}

        # 跳过表头，从第2行开始
        for row in self.worksheet.iter_rows(min_row=2, max_row=self.worksheet.max_row):
            product = self._parse_product_row(row)
            if product:
                self.products.append(product)
                # 用产品内编作为键
                if product.get('产品内编'):
                    self.products_dict[product['产品内编']] = product

    def _parse_product_row(self, row) -> Optional[Dict[str, str]]:
        """解析一行产品数据"""
        try:
            # 检查第一列是否为有效序号
            seq = row[self.COL_INDEX].value
            if seq is None or not str(seq).strip():
                return None

            product = {
                '序号': str(seq).strip() if seq else '',
                '产品内编': self._get_cell_value(row, self.COL_PRODUCT_CODE),
                'PI名建议': self._get_cell_value(row, self.COL_PI_NAME),
                '产品线': self._get_cell_value(row, self.COL_PRODUCT_LINE),
                '产品外观': self._get_cell_value(row, self.COL_APPEARANCE),
                '旧商品编码': self._get_cell_value(row, self.COL_OLD_CODE),
                '新商品编码': self._get_cell_value(row, self.COL_NEW_CODE),
                '报关名称': self._get_cell_value(row, self.COL_DECLARE_NAME),
                '运输鉴定名称': self._get_cell_value(row, self.COL_TRANSPORT_NAME),
                '用途': self._get_cell_value(row, self.COL_USAGE),
                '成分': self._get_cell_value(row, self.COL_COMPONENTS),
                '退税税率': self._get_cell_value(row, self.COL_TAX_REFUND),
                '监管条件': self._get_cell_value(row, self.COL_REGULATION),
                '检验检疫': self._get_cell_value(row, self.COL_INSPECTION),
            }

            # 如果产品内编为空，跳过
            if not product['产品内编']:
                return None

            return product

        except Exception:
            return None

    def _get_cell_value(self, row, col_index: int) -> str:
        """获取单元格值"""
        try:
            cell = row[col_index]
            value = cell.value
            if value is None:
                return ''
            return str(value).strip()
        except Exception:
            return ''

    def search_by_product_code(self, product_code: str) -> Optional[Dict[str, str]]:
        """
        通过产品内编搜索

        Args:
            product_code: 产品内编

        Returns:
            产品信息字典，未找到返回None
        """
        return self.products_dict.get(product_code)

    def search_by_name(self, name: str) -> List[Dict[str, str]]:
        """
        通过品名搜索（模糊匹配）

        Args:
            name: 品名（部分匹配）

        Returns:
            匹配的产品列表
        """
        name = name.lower()
        results = []

        for product in self.products:
            declare_name = product.get('报关名称', '').lower()
            transport_name = product.get('运输鉴定名称', '').lower()
            pi_name = product.get('PI名建议', '').lower()

            if name in declare_name or name in transport_name or name in pi_name:
                results.append(product)

        return results

    def get_all_products(self) -> List[Dict[str, str]]:
        """获取所有产品列表"""
        return self.products.copy()

    def get_count(self) -> int:
        """获取产品数量"""
        return len(self.products)

    def get_version(self) -> str:
        """获取版本信息（从文件名提取）"""
        if self.filepath:
            filename = os.path.basename(self.filepath)
            # 尝试从文件名提取日期
            import re
            match = re.search(r'(\d{4}\.\d{1,2}\.\d{1,2}|\d{4}-\d{2}-\d{2})', filename)
            if match:
                return match.group(1)
        return '未知'