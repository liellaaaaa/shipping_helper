"""
订舱单生成器
生成完整的订舱单文档
"""

import os
from typing import Optional, Dict

from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .data_models import ShipmentData


class BookingGenerator:
    """
    订舱单生成器
    生成完整的订舱单，包含收发货信息、货物信息、包装信息
    """

    def __init__(self):
        pass

    def generate(self, shipment: ShipmentData, output_dir: str = None) -> str:
        """
        生成订舱单

        Args:
            shipment: 完整货物信息
            output_dir: 输出目录

        Returns:
            生成的订舱单文件路径
        """
        if output_dir is None:
            output_dir = os.path.join(os.path.expanduser("~"), "Desktop")

        if not os.path.exists(output_dir):
            os.makedirs(output_dir)

        doc = Document()

        # 设置页面
        sections = doc.sections
        for section in sections:
            section.top_margin = Cm(2)
            section.bottom_margin = Cm(2)
            section.left_margin = Cm(2)
            section.right_margin = Cm(2)

        # 标题
        title = doc.add_heading("订舱单", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 副标题 - PI号
        if shipment.pi_no:
            pi_para = doc.add_paragraph()
            pi_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pi_run = pi_para.add_run(f"PI号: {shipment.pi_no}")
            pi_run.bold = True
            pi_run.font.size = Pt(12)

        doc.add_paragraph()  # 空行

        # 发货人信息
        self._add_section_title(doc, "发货人信息 (Shipper)")
        self._add_field(doc, "名称:", shipment.shipper or "")
        self._add_field(doc, "地址:", shipment.shipper_address or "")

        # 收货人信息
        self._add_section_title(doc, "收货人信息 (Consignee)")
        self._add_field(doc, "名称:", shipment.consignee or "")
        self._add_field(doc, "地址:", shipment.consignee_address or "")

        # 通知人信息
        self._add_section_title(doc, "通知人信息 (Notify Party)")
        self._add_field(doc, "名称:", shipment.notifier or "")
        self._add_field(doc, "地址:", shipment.notifier_address or "")

        # 货物信息
        self._add_section_title(doc, "货物信息 (Cargo Details)")

        # 产品名称
        product_name = shipment.product_name_cn or shipment.product_name_en or ""
        if shipment.cargo:
            if shipment.cargo.product_name_cn:
                product_name = shipment.cargo.product_name_cn
            if shipment.cargo.product_name_en:
                product_name += f" / {shipment.cargo.product_name_en}"

        self._add_field(doc, "品名:", product_name)
        self._add_field(doc, "内部编号:", shipment.internal_code or "")
        self._add_field(doc, "H.S. Code:", shipment.hs_code or "")

        if shipment.cargo:
            appearance = shipment.cargo.appearance_cn or shipment.cargo.appearance_en or ""
            self._add_field(doc, "外观与性状:", appearance)
            self._add_field(doc, "运输类型:", shipment.cargo.transport_type or "")
            if shipment.cargo.report_no:
                self._add_field(doc, "鉴定书编号:", shipment.cargo.report_no)

        # 包装信息
        self._add_section_title(doc, "包装信息 (Packing Details)")

        if shipment.package_info:
            pkg = shipment.package_info
            self._add_field(doc, "包装类型:", pkg.package_type or "")
            self._add_field(doc, "数量:", f"{pkg.quantity_kg} KG" if pkg.quantity_kg else "")
            self._add_field(doc, "桶数:", str(pkg.drums) if pkg.drums else "")
            self._add_field(doc, "卡板数:", str(pkg.pallets) if pkg.pallets else "")
            self._add_field(doc, "体积:", f"{pkg.volume_cbm} CBM" if pkg.volume_cbm else "")
            self._add_field(doc, "毛重:", f"{pkg.gross_weight_kg} KG" if pkg.gross_weight_kg else "")
            self._add_field(doc, "卡板规格:", pkg.pallet_type or "")
        else:
            for label in ["包装类型:", "数量:", "桶数:", "卡板数:", "体积:", "毛重:"]:
                self._add_field(doc, label, "")

        # 卸货港
        self._add_section_title(doc, "目的地 (Destination)")
        self._add_field(doc, "卸货港:", shipment.destination or "")

        # MSDS信息
        if shipment.msds and shipment.msds.physical_chemical:
            self._add_section_title(doc, "理化特性")
            pc = shipment.msds.physical_chemical
            self._add_field(doc, "pH值:", pc.ph or "")
            self._add_field(doc, "密度:", pc.density or "")
            self._add_field(doc, "沸点:", pc.boiling_point or "")
            self._add_field(doc, "闪点:", pc.flash_point or "")
            self._add_field(doc, "溶解性:", pc.solubility or "")

        # 生成文件名
        pi_no = shipment.pi_no or "Booking"
        product_name_file = shipment.product_name_cn or shipment.product_name_en or ""
        filename = f"订舱单-{pi_no}-{product_name_file}.docx"
        filename = filename.replace("/", "-").replace("\\", "-")
        output_path = os.path.join(output_dir, filename)

        doc.save(output_path)
        return output_path

    def _add_section_title(self, doc, title: str):
        """添加小节标题"""
        para = doc.add_paragraph()
        run = para.add_run(title)
        run.bold = True
        run.font.size = Pt(11)
        para.paragraph_format.space_after = Pt(3)

    def _add_field(self, doc, label: str, value: str):
        """添加字段行"""
        para = doc.add_paragraph()
        run = para.add_run(label)
        run.bold = True
        run.font.size = Pt(10)
        para.add_run(value)
        para.paragraph_format.space_after = Pt(2)