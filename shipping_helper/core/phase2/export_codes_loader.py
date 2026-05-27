"""
出口商品编码加载器
从 Excel 加载出口商品编码数据
"""

import os
from typing import List, Optional, Dict
from openpyxl import load_workbook

from .data_models import ExportCode


class ExportCodesLoader:
    """
    出口商品编码加载器
    文件路径：02.订舱出货/2024.12.5 最新出口商品编码及报关成分.xlsx
    """

    def __init__(self):
        self.filepath = ""
        self.data: List[ExportCode] = []
        self._loaded = False

    def load(self, filepath: str) -> bool:
        """
        加载Excel文件
        :param filepath: Excel文件路径
        :return: 加载是否成功
        """
        if not os.path.exists(filepath):
            return False

        self.filepath = filepath
        self.data = []
        self._loaded = True

        try:
            wb = load_workbook(filepath, data_only=True)
            ws = wb.active

            # 获取表头行（第一行）
            headers = [cell.value for cell in ws[1]]
            headers = [h.strip() if h else h for h in headers]

            # 映射列索引
            col_map = self._map_columns(headers)

            # 解析数据行
            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                try:
                    export_code = self._parse_row(row, col_map)
                    if export_code:
                        self.data.append(export_code)
                except Exception:
                    continue

            wb.close()
            return True

        except Exception as e:
            self.data = []
            self._loaded = False
            return False

    def _map_columns(self, headers: List[str]) -> Dict[str, int]:
        """映射列名到索引"""
        col_map = {}

        # 产品内编
        for name in ["产品内编", "内编", "内部编号", "产品编号"]:
            if name in headers:
                col_map["internal_code"] = headers.index(name)
                break

        # 产品外观
        for name in ["产品外观", "外观", "性状"]:
            if name in headers:
                col_map["appearance"] = headers.index(name)
                break

        # 新商品编码
        for name in ["新商品编码", "商品编码", "海关编码", "HS编码", "H.S.Code"]:
            if name in headers:
                col_map["hs_code"] = headers.index(name)
                break

        # 报关名称
        for name in ["报关名称", "品名", "名称"]:
            if name in headers:
                col_map["customs_name"] = headers.index(name)
                break

        # 成分
        for name in ["成分", "组成", "配方"]:
            if name in headers:
                col_map["composition"] = headers.index(name)
                break

        return col_map

    def _parse_row(self, row: tuple, col_map: Dict[str, int]) -> Optional[ExportCode]:
        """解析单行数据"""
        if not col_map:
            return None

        internal_code = ""
        appearance = ""
        hs_code = ""
        customs_name = ""
        composition = ""

        if "internal_code" in col_map:
            val = row[col_map["internal_code"]]
            internal_code = str(val).strip() if val else ""

        if "appearance" in col_map:
            val = row[col_map["appearance"]]
            appearance = str(val).strip() if val else ""

        if "hs_code" in col_map:
            val = row[col_map["hs_code"]]
            hs_code = str(val).strip() if val else ""

        if "customs_name" in col_map:
            val = row[col_map["customs_name"]]
            customs_name = str(val).strip() if val else ""

        if "composition" in col_map:
            val = row[col_map["composition"]]
            composition = str(val).strip() if val else ""

        # 至少需要有内部编号才认为是有效数据
        if not internal_code:
            return None

        return ExportCode(
            internal_code=internal_code,
            appearance=appearance,
            hs_code=hs_code,
            customs_name=customs_name,
            composition=composition
        )

    def find_by_internal_code(self, code: str) -> Optional[ExportCode]:
        """通过产品内编查找"""
        code = code.strip()
        for ec in self.data:
            if ec.internal_code == code:
                return ec
        return None

    def find_by_name(self, name: str) -> List[ExportCode]:
        """通过品名模糊搜索"""
        name = name.strip().lower()
        results = []
        for ec in self.data:
            if (name in ec.customs_name.lower() or
                name in ec.internal_code.lower() or
                name in ec.hs_code.lower()):
                results.append(ec)
        return results

    def get_all(self) -> List[ExportCode]:
        """获取所有编码"""
        return self.data

    def is_loaded(self) -> bool:
        """是否已加载"""
        return self._loaded

    def get_count(self) -> int:
        """获取数据条数"""
        return len(self.data)


if __name__ == "__main__":
    # 测试代码
    loader = ExportCodesLoader()
    test_path = r"c:\Users\windows\Desktop\shipping_helper\shipping_helper\02.订舱出货\2024.12.5 最新出口商品编码及报关成分.xlsx"

    if loader.load(test_path):
        print(f"加载成功，共 {loader.get_count()} 条数据")
        # 测试查找
        result = loader.find_by_internal_code("1011")
        if result:
            print(f"查到: {result.customs_name}, HS: {result.hs_code}")
    else:
        print("加载失败")